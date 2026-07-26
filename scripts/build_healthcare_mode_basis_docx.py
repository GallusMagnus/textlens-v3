from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "healthcare-mode-basis.md"
OUTPUT = ROOT / "docs" / "healthcare-mode-basis.docx"


def set_spacing(style, before=0, after=6, line=1.1):
    fmt = style.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def set_font(style, name="Calibri", size=11, color="000000", bold=False):
    font = style.font
    font.name = name
    font.size = Pt(size)
    font.color.rgb = RGBColor.from_string(color)
    font.bold = bold


def set_list_indents(style, left=0.5, first=-0.25, after=6, line=1.167):
    fmt = style.paragraph_format
    fmt.left_indent = Inches(left)
    fmt.first_line_indent = Inches(first)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(10)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(table, top=80, start=120, bottom=80, end=120):
    tbl_pr = table._tbl.tblPr
    margins = tbl_pr.first_child_found_in("w:tblCellMar")
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for margin_name, value in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():
        node = margins.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_text_with_code_runs(paragraph, text, bold=False):
    parts = re.split(r"(`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        run = paragraph.add_run(part[1:-1] if part.startswith("`") and part.endswith("`") else part)
        run.bold = bold
        run.font.name = "Consolas" if part.startswith("`") and part.endswith("`") else "Calibri"
        run.font.size = Pt(10 if part.startswith("`") and part.endswith("`") else 11)


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    add_text_with_code_runs(p, text)


def add_numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    add_text_with_code_runs(p, text)


def build_doc():
    doc = Document()

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    set_font(styles["Normal"], size=11)
    set_spacing(styles["Normal"], after=6, line=1.1)

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        set_font(styles[name], size=size, color=color, bold=True)
        set_spacing(styles[name], before=before, after=after, line=1.1)

    set_list_indents(styles["List Bullet"], after=6)
    set_list_indents(styles["List Bullet 2"], left=0.75, first=-0.25, after=4)
    set_list_indents(styles["List Number"], after=6)

    title_style = styles.add_style("TextLens Title", 1)
    set_font(title_style, size=20, color="0B2545", bold=True)
    set_spacing(title_style, before=0, after=4, line=1.1)

    subtitle_style = styles.add_style("TextLens Subtitle", 1)
    set_font(subtitle_style, size=11, color="5B677A")
    set_spacing(subtitle_style, before=0, after=14, line=1.1)

    callout_style = styles.add_style("TextLens Callout", 1)
    set_font(callout_style, size=10, color="1F3A5F")
    set_spacing(callout_style, before=4, after=10, line=1.1)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer.add_run("TextLens Healthcare Mode Basis")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor.from_string("7B8794")

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    title = lines[0].lstrip("# ").strip()
    doc.add_paragraph(title, style="TextLens Title")
    doc.add_paragraph("TextLens method governance note", style="TextLens Subtitle")

    body_lines = lines[1:]
    in_source_families = False
    numbered_buffer = []

    def flush_numbered_buffer():
        nonlocal numbered_buffer
        for item in numbered_buffer:
            add_numbered(doc, item)
        numbered_buffer = []

    for raw_line in body_lines:
        line = raw_line.rstrip()
        if not line:
            continue

        if line.startswith("## "):
            flush_numbered_buffer()
            heading = line[3:].strip()
            in_source_families = heading == "Source Families"
            doc.add_heading(heading, level=1)
            continue

        numbered_match = re.match(r"^(\d+)\.\s+(.*)$", line)
        if numbered_match:
            if in_source_families:
                flush_numbered_buffer()
                doc.add_heading(numbered_match.group(2), level=2)
            else:
                numbered_buffer.append(numbered_match.group(2))
            continue

        bullet_match = re.match(r"^-\s+(.*)$", line)
        nested_bullet_match = re.match(r"^\s+-\s+(.*)$", line)
        if bullet_match:
            flush_numbered_buffer()
            add_bullet(doc, bullet_match.group(1), level=0)
            continue
        if nested_bullet_match:
            flush_numbered_buffer()
            add_bullet(doc, nested_bullet_match.group(1), level=1)
            continue

        flush_numbered_buffer()
        p = doc.add_paragraph()
        add_text_with_code_runs(p, line)

    flush_numbered_buffer()

    info_table = doc.add_table(rows=2, cols=2)
    info_table.style = "Table Grid"
    info_table.autofit = False
    set_cell_margins(info_table)
    widths = [Inches(1.8), Inches(4.7)]
    rows = [
        ("Markdown source", "docs/healthcare-mode-basis.md"),
        ("Implementation anchor", "src/analysis/policies/modePolicies.ts"),
    ]
    for row_idx, row in enumerate(info_table.rows):
        for cell_idx, cell in enumerate(row.cells):
            cell.width = widths[cell_idx]
            set_cell_text(cell, rows[row_idx][cell_idx], bold=cell_idx == 0)
            if cell_idx == 0:
                shade_cell(cell, "F2F4F7")

    doc.save(OUTPUT)


if __name__ == "__main__":
    build_doc()
    print(OUTPUT)
